from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from threading import Lock

from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

BASE_DIR = Path(__file__).parent
OUTPUT_DIR = BASE_DIR / "output"
PLACEHOLDER_PATTERN = re.compile(r"\{\{([a-zA-Z_][a-zA-Z0-9_]*)\}\}")

app = FastAPI(title="Figma PDF Generator")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)

_rows: list[dict[str, object]] = []
_lock = Lock()


def read_excel(source: Path | bytes) -> list[dict[str, object]]:
    try:
        if isinstance(source, Path):
            if not source.exists():
                raise ValueError(f"Input file not found: {source.name}")
            if source.suffix.lower() != ".xlsx":
                raise ValueError("Input file must have .xlsx extension")
        workbook_source = io.BytesIO(source) if isinstance(source, bytes) else source
        workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    except Exception as exc:
            source_name = source.name if isinstance(source, Path) else "uploaded Excel file"
            raise ValueError(f"Cannot open {source_name}: {exc}") from exc

    try:
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not values:
        raise ValueError("Excel file has no header row")
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    if not headers or any(not header for header in headers):
        raise ValueError("Excel header must not contain empty column names")
    if len(headers) != len(set(headers)):
        raise ValueError("Excel header columns must be unique")

    rows = []
    for row in values[1:]:
        if not any(value is not None for value in row):
            continue
        item = dict(zip(headers, row))
        rows.append(item)
    if not rows:
        raise ValueError("Excel file has no data rows")
    return rows


@app.get("/rows")
def rows() -> dict[str, object]:
    with _lock:
        if not _rows:
            raise HTTPException(status_code=400, detail="Upload an .xlsx file in the plugin first")
        return {"count": len(_rows), "rows": _rows}


@app.post("/excel")
async def upload_excel(request: Request) -> dict[str, int]:
    filename = request.headers.get("x-filename", "")
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload an .xlsx file")
    content = await request.body()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded Excel file is empty")
    try:
        loaded_rows = read_excel(content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    with _lock:
        _rows.clear()
        _rows.extend(loaded_rows)
    return {"count": len(_rows)}


@app.post("/pdf")
async def save_pdf(request: Request) -> dict[str, str]:
    payload = await request.json()
    filename = payload.get("filename", "")
    if not re.fullmatch(r"[^\\/:*?\"<>|]+_КП\.pdf", filename):
        raise HTTPException(status_code=400, detail="filename must look like <uni_name>_КП.pdf")
    parts = payload.get("parts", [])
    if len(parts) != 2:
        raise HTTPException(status_code=400, detail="Exactly two PDF parts are required")
    OUTPUT_DIR.mkdir(exist_ok=True)
    writer = PdfWriter()
    try:
        for part in parts:
            content = bytes(part)
            if not content.startswith(b"%PDF"):
                raise ValueError("Request contains a non-PDF part")
            for page in PdfReader(io.BytesIO(content)).pages:
                writer.add_page(page)
        with (OUTPUT_DIR / filename).open("wb") as output:
            writer.write(output)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not merge PDF parts: {exc}") from exc
    total = len(_rows) or "?"
    print(f"[{filename[:3]}/{total}] OK", flush=True)
    return {"filename": filename}


@app.post("/reset")
def reset_output() -> dict[str, str]:
    OUTPUT_DIR.mkdir(exist_ok=True)
    for pdf in OUTPUT_DIR.glob("*.pdf"):
        pdf.unlink()
    return {"status": "ok"}


@app.post("/zip")
def zip_output() -> dict[str, str]:
    archive = create_zip()
    print(f"Generation completed. {len(list(OUTPUT_DIR.glob('*.pdf')))} PDFs generated.", flush=True)
    return {"filename": archive.name}


def create_zip() -> Path:
    archive = BASE_DIR / "output.zip"
    pdfs = sorted(OUTPUT_DIR.glob("*.pdf"))
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for pdf in pdfs:
            zip_file.write(pdf, pdf.name)
    return archive


if __name__ == "__main__":
    import uvicorn

    print("Figma PDF bridge listening at http://127.0.0.1:8000")
    print("Keep this process running while the Figma plugin generates PDFs.")
    uvicorn.run(app, host="127.0.0.1", port=8000)